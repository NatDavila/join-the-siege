import pytesseract
from pdf2image import convert_from_bytes
from PIL import Image
import cv2
import numpy as np
from io import BytesIO
import docx
import openpyxl
import xlrd
import logging

pil_image = Image.Image

def preprocess_image(image: pil_image) -> pil_image:
    """
    Convert image to grayscale and apply thresholding to improve OCR accuracy

    parameters:
        image: PIL Image object

    returns:
        pil_image: The preprocessed image
    """
    try:
        gray = cv2.cvtColor(np.array(image), cv2.COLOR_RGB2GRAY)
        _, thresh = cv2.threshold(gray, 150, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        return Image.fromarray(thresh)
    except Exception as e:
        logging.error(f"Error in preprocessing image: {e}")
        raise


def ocr_image(image_file: bytes) -> str:
    """ 
    Extract text from an image (JPG, PNG, JPEG) 

    parameters:
        image_file: The byte string containing the image file 

    returns:
        text: The extracted text from the image
    """
    image = Image.open(BytesIO(image_file)) 
    preprocessed_image = preprocess_image(image)
    text = pytesseract.image_to_string(preprocessed_image)
    return text


def ocr_pdf(pdf_bytes: bytes) -> str:
    """
    Extract text from a PDF by converting it to images first 

    parameters:
        pdf_bytes: The byte string containing the PDF file

    returns:
        text: The extracted text from the PDF
    """
    images = convert_from_bytes(pdf_bytes)
    return "\n".join([pytesseract.image_to_string(preprocess_image(img)) for img in images])


def ocr_docx(docx_bytes: bytes) -> str:
    """
    Extract text from a docx file

    parameters: 
        docx_bytes: The byte string containing the docx file
    
    returns:
        text: The extracted text from the docx file
    """
    with BytesIO(docx_bytes) as temp_file:
        doc = docx.Document(temp_file)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text


def ocr_excel(excel_bytes: bytes, file_extension: str) -> str:
    """
    Extract text from both XLSX and XLS files based on extension

    parameters:
        excel_bytes: The byte string containing the excel file
        file_extension: The file extension (XLSX or XLS)

    returns:
        text: The extracted text from the excel file
    """
    text = ""
    
    # If file is .xlsx, use openpyxl
    if file_extension == 'xlsx':
        with BytesIO(excel_bytes) as temp_file:
            workbook = openpyxl.load_workbook(temp_file, data_only=True)
            for sheet in workbook:
                for row in sheet.iter_rows(values_only=True):
                    row_text = ""
                    for cell in row:
                        if cell is not None:
                            row_text += str(cell) + " "
                    text += row_text + "\n"
    
    # If file is .xls, use xlrd
    elif file_extension == 'xls':
        with BytesIO(excel_bytes) as temp_file:
            workbook = xlrd.open_workbook(file_contents=temp_file.read())
            for sheet in workbook.sheets():
                for row_idx in range(sheet.nrows):
                    row = sheet.row(row_idx)
                    row_text = ""
                    for cell in row:
                        if cell.value:
                            row_text += str(cell.value) + " "
                    text += row_text + "\n"
    return text


def text_extractor(file_bytes: bytes, file_extension: str) -> str:
    """
    Determine the appropriate OCR function based on file type and extract text

    parameters:
        file_bytes: The byte string containing the file
        file_extension: The file extension

    returns:
        text: The extracted text from the file
    """
    ocr_functions = {
        'jpg': ocr_image,
        'jpeg': ocr_image,
        'png': ocr_image,
        'pdf': ocr_pdf,
        'docx': ocr_docx,
        'xlsx': ocr_excel,
        'xls': ocr_excel,
        'txt': lambda b: b.decode('utf-8')
    }

    try:
        if file_extension not in ocr_functions:
            raise ValueError(f"Unsupported file format: {file_extension}")

        return ocr_functions[file_extension](file_bytes)

    except Exception as e:
        logging.error(f"Error processing file: {e}")
        raise
    